from copy import copy
from pathlib import Path

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]
TEMPLATE = ROOT / "web-import-template-v0.1.xlsx"


def copy_header_style(worksheet, source_column: int, target_columns: list[int]) -> None:
    source = worksheet.cell(1, source_column)
    for column in target_columns:
        target = worksheet.cell(1, column)
        target.font = copy(source.font)
        target.fill = copy(source.fill)
        target.border = copy(source.border)
        target.alignment = copy(source.alignment)
        target.number_format = source.number_format
        target.protection = copy(source.protection)


def append_unique(worksheet, key_columns: tuple[int, ...], rows: list[tuple]) -> None:
    existing = {
        tuple(worksheet.cell(row, column).value for column in key_columns)
        for row in range(2, worksheet.max_row + 1)
    }
    for values in rows:
        key = tuple(values[column - 1] for column in key_columns)
        if key not in existing:
            worksheet.append(values)


def main() -> None:
    workbook = load_workbook(TEMPLATE)

    parks = workbook["园区档案"]
    if parks.cell(1, 4).value != "longitude":
        parks.insert_cols(4, 2)
        parks.cell(1, 4, "longitude")
        parks.cell(1, 5, "latitude")
        copy_header_style(parks, 6, [4, 5])

    stores = workbook["门店"]
    if stores.cell(1, 4).value != "park_id":
        stores.insert_cols(4, 3)
        stores.cell(1, 4, "park_id")
        stores.cell(1, 5, "channel_type")
        stores.cell(1, 6, "reporting_authorized")
        copy_header_style(stores, 7, [4, 5, 6])
    headers = [cell.value for cell in stores[1]]
    if "city" not in headers:
        longitude_column = headers.index("longitude") + 1
        stores.insert_cols(longitude_column, 1)
        stores.cell(1, longitude_column, "city")
        copy_header_style(stores, longitude_column + 1, [longitude_column])

    dictionary = workbook["字段字典"]
    append_unique(
        dictionary,
        (1, 2),
        [
            ("园区档案", "longitude", "number", "否", "EPSG:4326", "园区经度", 125.182),
            ("园区档案", "latitude", "number", "否", "EPSG:4326", "园区纬度", 44.432),
            ("门店", "park_id", "string", "否", None, "所属园区编号；有效新门店必须填写", "PARK-001"),
            ("门店", "channel_type", "enum", "否", "TRADITIONAL_STORE/THIRD_SPACE", "渠道分类；有效新门店必须填写", "THIRD_SPACE"),
            ("门店", "reporting_authorized", "boolean", "否", None, "是否允许 B02 上报经营日报", True),
            ("门店", "city", "string", "否", None, "公开大屏城市标签", "长春"),
        ],
    )
    enums = workbook["枚举字典"]
    append_unique(
        enums,
        (1, 2),
        [
            ("channel_type", "TRADITIONAL_STORE", "传统门店", "传统零售终端"),
            ("channel_type", "THIRD_SPACE", "第三空间", "餐饮、会客厅、主题驿站等体验式终端"),
        ],
    )
    workbook.save(TEMPLATE)


if __name__ == "__main__":
    main()
